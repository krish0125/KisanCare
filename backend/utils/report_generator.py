"""
backend/utils/report_generator.py
---------------------------------
Generates PDF and Excel reports for crop cycle expenses and profits.
Uses reportlab (PDF) and openpyxl (Excel).
"""

import io
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
import openpyxl

def generate_pdf_report(cycle_doc: dict, summary: dict) -> bytes:
    """Generate a PDF report for a crop cycle."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    elements = []
    
    styles = getSampleStyleSheet()
    title_style = styles['Heading1']
    normal_style = styles['Normal']
    bold_style = ParagraphStyle('Bold', parent=styles['Normal'], fontName='Helvetica-Bold')
    
    # Title
    elements.append(Paragraph(f"Crop Cycle Report: {cycle_doc.get('cycle_name', 'Untitled')}", title_style))
    elements.append(Spacer(1, 12))
    
    # Metadata
    elements.append(Paragraph(f"Crop: {cycle_doc.get('crop', 'N/A').title()}", normal_style))
    elements.append(Paragraph(f"Land Area: {cycle_doc.get('land_area_acres', 0)} acres", normal_style))
    elements.append(Paragraph(f"Start Date: {cycle_doc.get('start_date', 'N/A')}", normal_style))
    elements.append(Spacer(1, 12))
    
    # Financial Summary Table
    elements.append(Paragraph("Financial Summary", styles['Heading2']))
    summary_data = [
        ["Total Investment", f"Rs. {summary.get('total_investment_inr', 0):,.2f}"],
        ["Expected Revenue", f"Rs. {summary.get('expected_revenue_inr', 0):,.2f}"],
        ["Actual Revenue", f"Rs. {summary.get('actual_revenue_inr', 0):,.2f}"],
        ["Current Revenue Est.", f"Rs. {summary.get('current_revenue_estimate_inr', 0):,.2f}"],
        ["Profit/Loss", f"Rs. {summary.get('profit_loss_inr', 0):,.2f}"],
        ["Profit/Loss per Acre", f"Rs. {summary.get('profit_loss_per_acre_inr', 0):,.2f}"]
    ]
    
    t_summary = Table(summary_data, colWidths=[200, 200])
    t_summary.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.white),
        ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
    ]))
    elements.append(t_summary)
    elements.append(Spacer(1, 20))
    
    # Expenses Table
    elements.append(Paragraph("Expense Breakdown", styles['Heading2']))
    
    expenses = cycle_doc.get('expenses', [])
    if not expenses:
        elements.append(Paragraph("No expenses recorded for this cycle.", normal_style))
    else:
        exp_data = [["Date", "Category", "Amount (Rs.)", "Note"]]
        for exp in expenses:
            exp_data.append([
                exp.get('date', ''),
                exp.get('category', '').title(),
                f"{exp.get('amount_inr', 0):,.2f}",
                exp.get('note', '')
            ])
            
        t_exp = Table(exp_data, colWidths=[80, 100, 100, 200])
        t_exp.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.black),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('ALIGN', (2, 1), (2, -1), 'RIGHT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.white),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ]))
        elements.append(t_exp)
        
    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()

def generate_excel_report(cycle_doc: dict, summary: dict) -> bytes:
    """Generate an Excel report for a crop cycle."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Crop Cycle Report"
    
    # Title & Metadata
    ws['A1'] = "Crop Cycle Report"
    ws['A1'].font = openpyxl.styles.Font(bold=True, size=14)
    
    ws['A3'] = "Cycle Name:"
    ws['B3'] = cycle_doc.get('cycle_name', 'Untitled')
    ws['A4'] = "Crop:"
    ws['B4'] = cycle_doc.get('crop', 'N/A').title()
    ws['A5'] = "Land Area (acres):"
    ws['B5'] = cycle_doc.get('land_area_acres', 0)
    ws['A6'] = "Start Date:"
    ws['B6'] = cycle_doc.get('start_date', 'N/A')
    
    # Summary
    ws['A8'] = "Financial Summary"
    ws['A8'].font = openpyxl.styles.Font(bold=True)
    
    summary_data = [
        ("Total Investment", summary.get('total_investment_inr', 0)),
        ("Expected Revenue", summary.get('expected_revenue_inr', 0)),
        ("Actual Revenue", summary.get('actual_revenue_inr', 0)),
        ("Current Revenue Est.", summary.get('current_revenue_estimate_inr', 0)),
        ("Profit/Loss", summary.get('profit_loss_inr', 0)),
        ("Profit/Loss per Acre", summary.get('profit_loss_per_acre_inr', 0))
    ]
    
    row = 9
    for key, val in summary_data:
        ws.cell(row=row, column=1, value=key).font = openpyxl.styles.Font(bold=True)
        ws.cell(row=row, column=2, value=val).number_format = '#,##0.00'
        row += 1
        
    # Expenses
    row += 2
    ws.cell(row=row, column=1, value="Expense Breakdown").font = openpyxl.styles.Font(bold=True)
    row += 1
    
    headers = ["Date", "Category", "Amount (Rs.)", "Note"]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.font = openpyxl.styles.Font(bold=True)
        
    row += 1
    expenses = cycle_doc.get('expenses', [])
    for exp in expenses:
        ws.cell(row=row, column=1, value=exp.get('date', ''))
        ws.cell(row=row, column=2, value=exp.get('category', '').title())
        ws.cell(row=row, column=3, value=exp.get('amount_inr', 0)).number_format = '#,##0.00'
        ws.cell(row=row, column=4, value=exp.get('note', ''))
        row += 1
        
    # Adjust column widths
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 40
    
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
